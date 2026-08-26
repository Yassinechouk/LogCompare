# -*- coding: utf-8 -*-
"""
GUI.py – Tkinter Application for LOGCOMPARE
============================================

Provides the :class:`LogCompareApp` window which lets users:

  - Select a reference ``.log`` file and a folder of comparison logs.
  - Build and colour-code a step-duration comparison table.
  - Export results to CSV and Excel (with an embedded duration chart).
  - Filter steps by name and toggle delta columns.
  - Persist session preferences between launches.
"""

# ── Standard library ──────────────────────────────────────────────────────
import csv
import time
import traceback
from collections import Counter
from pathlib import Path
from typing import Any, List, Optional

# ── Third-party ───────────────────────────────────────────────────────────
import matplotlib.pyplot as plt
import openpyxl
import tkinter as tk
from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg, NavigationToolbar2Tk  # pyright: ignore[reportPrivateImportUsage]
from matplotlib.figure import Figure
from openpyxl.chart import ScatterChart, Reference
from openpyxl.chart.axis import ChartLines, TextAxis
from openpyxl.chart.series import Series
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils import get_column_letter as _gcl
from tkinter import filedialog, messagebox, ttk
from tksheet import Sheet
from typing import Any as _Any

# ── Internal ──────────────────────────────────────────────────────────────
from constants import (
    AFFICHER_DELTAS_DEFAULT,
    APP_TITLE,
    COLOR_BAD,
    COLOR_DELTA_BG_DEFAULT,
    COLOR_EMPTY,
    COLOR_EQUAL_DARK,
    COLOR_NEAR_LIGHT,
    COLOR_REF_BG,
    COLOR_STEP_COLUMN_BG,
    COLOR_TOTAL_BG,
    COLOR_WARN,
    DEFAULT_THRESH_NEAR,
    DEFAULT_THRESH_YELLOW,
    INCLURE_INDEX_DEFAULT,
    MANDATORY_COLUMNS,
    SIGNATURE_ENABLED,
    SIGNATURE_TEXT,
)
from models import Etape
from parsing import affecter_blocs, extract_duree_test_seconds, parse_etapes
from pivot import construire_pivot
from prefs import load_prefs, save_prefs

class LogCompareApp(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title(APP_TITLE)
        self.geometry("1450x860")
        self.minsize(1180, 680)
        self.reference_path: Optional[Path] = None
        self.reference_etapes: List[Etape] = []
        self.logs_folder: Optional[Path] = None
        self.output_folder: Optional[Path] = None
        self.pivot_headers: List[str] = []
        self.pivot_rows: List[List[Any]] = []
        self.pivot_file: Optional[Path] = None
        self.hidden_cols: set[int] = set()
        self.file_label_map_excel: dict[str, str] = {}

        prefs = load_prefs()
        thr_near_saved = prefs.get("thr_near", prefs.get("thr_green", 0.0))
        thr_yellow_saved = prefs.get("thr_yellow", 0.0)
        self.var_show_deltas = tk.BooleanVar(value=prefs.get("show_deltas", AFFICHER_DELTAS_DEFAULT))
        self.var_thr_near = tk.DoubleVar(value=thr_near_saved)
        self.var_thr_yellow = tk.DoubleVar(value=thr_yellow_saved)
        self.var_search = tk.StringVar()

        try:
            self.Sheet = Sheet
            self.sheet_available = True
        except Exception:
            self.Sheet = None
            self.sheet_available = False

        try:
            self.openpyxl_available = True
        except Exception:
            self.openpyxl_available = False

        try:
            self.matplotlib_available = True
        except Exception:
            self.matplotlib_available = False

        self._build_ui()
        self._build_menu()
        self._bind_easter_egg()

        # Restore paths
        if "ref_path" in prefs:
            refp = Path(prefs["ref_path"])
            if refp.is_file():
                self.entry_ref.insert(0, str(refp))
                self._load_reference_file(refp)
        if "logs_folder" in prefs:
            fold = Path(prefs["logs_folder"])
            if fold.is_dir():
                self.entry_logs.insert(0, str(fold))
                self.logs_folder = fold
        if "output_folder" in prefs:
            outd = Path(prefs["output_folder"])
            if outd.is_dir():
                self.entry_out.insert(0, str(outd))
                self.output_folder = outd

        self.after(250, self._auto_apply_search)

    # ----- Menu & Signature -----
    def _build_menu(self):
        menubar = tk.Menu(self)
        if SIGNATURE_ENABLED:
            help_menu = tk.Menu(menubar, tearoff=False)
            help_menu.add_command(label="À propos", command=self._show_about)
            menubar.add_cascade(label="Aide", menu=help_menu)
        self.config(menu=menubar)

    def _show_about(self):
        txt = f"{APP_TITLE}\n\nComparaison de logs (premier bloc).\n"
        if SIGNATURE_ENABLED:
            txt += f"\nSignature : {SIGNATURE_TEXT}"
        messagebox.showinfo("À propos", txt)

    def _bind_easter_egg(self):
        if SIGNATURE_ENABLED:
            self.bind_all("<Control-Alt-l>", lambda e: self._show_signature_popup())

    def _show_signature_popup(self):
        top = tk.Toplevel(self)
        top.title("Signature")
        msg = SIGNATURE_TEXT if SIGNATURE_ENABLED else "Signature désactivée"
        ttk.Label(top, text=msg, foreground="#666", padding=12, wraplength=320, justify="center").pack()
        ttk.Button(top, text="Fermer", command=top.destroy).pack(pady=6)

    # ----- UI principal -----
    def _build_ui(self):
        frm_paths = ttk.LabelFrame(self, text="Chemins", padding=8)
        frm_paths.pack(fill="x", padx=8, pady=6)

        ttk.Label(frm_paths, text="Fichier référence:").grid(row=0, column=0, sticky="w")
        self.entry_ref = ttk.Entry(frm_paths, width=95)
        self.entry_ref.grid(row=0, column=1, sticky="we", padx=4)
        ttk.Button(frm_paths, text="Parcourir", command=self.choose_reference).grid(row=0, column=2, padx=4)

        ttk.Label(frm_paths, text="Dossier logs:").grid(row=1, column=0, sticky="w")
        self.entry_logs = ttk.Entry(frm_paths, width=95)
        self.entry_logs.grid(row=1, column=1, sticky="we", padx=4)
        ttk.Button(frm_paths, text="Parcourir", command=self.choose_logs_folder).grid(row=1, column=2, padx=4)

        ttk.Label(frm_paths, text="Dossier sortie (optionnel):").grid(row=2, column=0, sticky="w")
        self.entry_out = ttk.Entry(frm_paths, width=95)
        self.entry_out.grid(row=2, column=1, sticky="we", padx=4)
        ttk.Button(frm_paths, text="Parcourir", command=self.choose_output_folder).grid(row=2, column=2, padx=4)

        frm_paths.columnconfigure(1, weight=1)

        frm_opts = ttk.LabelFrame(self, text="Options / Filtres / Export", padding=6)
        frm_opts.pack(fill="x", padx=8, pady=4)

        ttk.Checkbutton(frm_opts, text="Afficher deltas", variable=self.var_show_deltas).pack(side="left", padx=10)

        ttk.Label(frm_opts, text="Seuil proche % ≤").pack(side="left", padx=(35, 4))
        tk.Entry(frm_opts, width=6, textvariable=self.var_thr_near).pack(side="left")
        ttk.Label(frm_opts, text="Seuil jaune % ≤").pack(side="left", padx=(14, 4))
        tk.Entry(frm_opts, width=6, textvariable=self.var_thr_yellow).pack(side="left")

        ttk.Label(frm_opts, text="Recherche étape:").pack(side="left", padx=(30, 4))
        search_entry = ttk.Entry(frm_opts, width=26, textvariable=self.var_search)
        search_entry.pack(side="left")
        search_entry.bind("<KeyRelease>", lambda e: self.filter_rows())

        ttk.Button(frm_opts, text="Construire tableau", command=self.build_table).pack(side="left", padx=14)
        ttk.Button(frm_opts, text="Recolorer", command=self.recolor_table).pack(side="left")

        self.btn_export_xlsx = ttk.Button(
            frm_opts,
            text="Extraire Excel",
            command=self.export_excel,
            state="normal" if self.openpyxl_available else "disabled"
        )
        self.btn_export_xlsx.pack(side="left", padx=14)

        self.btn_curve = ttk.Button(
            frm_opts,
            text="Courbe durée test",
            command=self.show_duration_curve,
            state="normal" if self.matplotlib_available else "disabled"
        )
        self.btn_curve.pack(side="left", padx=6)

        ttk.Button(frm_opts, text="Colonnes...", command=self.manage_columns).pack(side="left", padx=6)

        if not self.sheet_available:
            ttk.Label(self, foreground="red",
                      text="tksheet non installé : pip install tksheet").pack(fill="x", padx=8, pady=4)
        if not self.openpyxl_available:
            ttk.Label(self, foreground="orange",
                      text="Export Excel indisponible (pip install openpyxl)").pack(fill="x", padx=8, pady=2)
        if not self.matplotlib_available:
            ttk.Label(self, foreground="orange",
                      text="Courbe indisponible (pip install matplotlib)").pack(fill="x", padx=8, pady=2)

        frm_table = ttk.LabelFrame(self, text="Tableau comparatif (1er bloc)", padding=4)
        frm_table.pack(fill="both", expand=True, padx=8, pady=4)

        if self.sheet_available:
            self.sheet = Sheet(frm_table, data=[], headers=[], height=520, width=1400)
            self.sheet.enable_bindings((
                "single_select", "row_select", "column_select", "drag_select",
                "arrowkeys", "row_height_resize", "column_width_resize",
                "copy", "rc_select", "hide_columns", "hide_rows", "undo", "edit_cell"
            )) # pyright: ignore[reportArgumentType]
            self.sheet.pack(fill="both", expand=True)
        else:
            self.sheet = None

        frm_log = ttk.LabelFrame(self, text="Journal", padding=6)
        frm_log.pack(fill="both", expand=False, padx=8, pady=4)
        self.txt_log = tk.Text(frm_log, height=9, wrap="word")
        self.txt_log.pack(fill="both", expand=True)

        status_frame = ttk.Frame(self)
        status_frame.pack(fill="x", side="bottom")
        self.status_var = tk.StringVar(value="Prêt.")
        ttk.Label(status_frame, textvariable=self.status_var, anchor="w").pack(fill="x", side="left", expand=True)
        if SIGNATURE_ENABLED:
            ttk.Label(status_frame, text=SIGNATURE_TEXT, anchor="e", foreground="#888888",
                      font=("Segoe UI", 8)).pack(side="right", padx=8, pady=2)

        footer = ttk.Frame(self)
        footer.pack(fill="x", side="bottom")
        ttk.Label(footer, text="aphordite", anchor="center",
                  foreground="#888888", font=("Segoe UI", 9, "italic")).pack(pady=2)

    def log(self, msg: str):
        self.txt_log.insert("end", msg + "\n")
        self.txt_log.see("end")
        self.status_var.set(msg)
        print(msg)

    def choose_reference(self):
        p = filedialog.askopenfilename(title="Fichier log de référence",
                                       filetypes=[("Fichiers log", "*.log"), ("Tous", "*.*")])
        if not p:
            return
        path = Path(p)
        self.entry_ref.delete(0, "end")
        self.entry_ref.insert(0, str(path))
        self._load_reference_file(path)
        self._save_prefs_state()

    def _load_reference_file(self, path: Path):
        if not path.is_file():
            messagebox.showerror("Erreur", "Fichier référence invalide.")
            return
        try:
            etapes = parse_etapes(path, manage_rollover=True)
            affecter_blocs(etapes)
            self.reference_etapes = etapes
            self.reference_path = path
            self.log(f"Référence chargée : {path.name} (étapes={len(etapes)})")
        except Exception:
            self.log("Erreur chargement référence:\n" + traceback.format_exc())
            messagebox.showerror("Erreur", "Impossible de charger le fichier de référence.")

    def choose_logs_folder(self):
        d = filedialog.askdirectory(title="Dossier logs")
        if not d:
            return
        self.entry_logs.delete(0, "end")
        self.entry_logs.insert(0, d)
        self.logs_folder = Path(d)
        self.log(f"Dossier logs sélectionné : {self.logs_folder}")
        self._save_prefs_state()

    def choose_output_folder(self):
        d = filedialog.askdirectory(title="Dossier sortie CSV/Excel")
        if not d:
            return
        self.entry_out.delete(0, "end")
        self.entry_out.insert(0, d)
        self.output_folder = Path(d)
        self.log(f"Dossier sortie sélectionné : {self.output_folder}")
        self._save_prefs_state()

    def build_table(self):
        if not self.reference_path:
            messagebox.showwarning("Référence", "Sélectionne un fichier de référence.")
            return
        logs_dir = self.entry_logs.get().strip()
        logs_folder = Path(logs_dir) if logs_dir else self.reference_path.parent
        out_dir_raw = self.entry_out.get().strip()
        out_dir = Path(out_dir_raw) if out_dir_raw else None
        try:
            autres = sorted([p for p in logs_folder.glob("*.log")
                             if self.reference_path and p.resolve() != self.reference_path.resolve()])
            self.file_label_map_excel = {f"Fichier_{i+1}": p.name for i, p in enumerate(autres)}
        except Exception:
            self.file_label_map_excel = {}
        try:
            headers, rows, pivot_file = construire_pivot(
                reference_path=self.reference_path,
                reference_etapes=self.reference_etapes,
                dossier_logs=logs_folder,
                include_index=False,
                include_deltas=self.var_show_deltas.get(),
                output_dir=out_dir
            )
            self.pivot_headers = headers
            self.pivot_rows = rows
            self.pivot_file = pivot_file
            self.hidden_cols.clear()
            self.update_sheet()
            self.log(f"Tableau construit & CSV généré : {pivot_file.name}")
            self._save_prefs_state()
        except Exception:
            self.log("Erreur construction pivot:\n" + traceback.format_exc())
            messagebox.showerror("Erreur", "Impossible de construire le tableau.")

    def show_duration_curve(self):
        if not self.pivot_rows:
            messagebox.showwarning("Données", "Veuillez construire le tableau.")
            return
        first_row = self.pivot_rows[0]
        durations = [v for v in first_row if isinstance(v, int)]
        counts = {}
        for d in durations:
            counts[d] = counts.get(d, 0) + 1
        ys = sorted(counts.keys())
        xs = [counts[y] for y in ys]
        sorted_pairs = sorted(zip(xs, ys))
        xs_sorted, ys_sorted = zip(*sorted_pairs)
        ref_s: Optional[int] = None
        if self.reference_path:
            try:
                v = extract_duree_test_seconds(self.reference_path)
                ref_s = v if isinstance(v, int) else None
            except Exception:
                ref_s = None
        plt.figure(figsize=(8, 4))
        plt.plot(xs_sorted, ys_sorted, marker="o", linestyle="-")
        if isinstance(ref_s, int):
            plt.axhline(ref_s, color="#d62728", linestyle="--", linewidth=1.5, label=f"Référence: {ref_s} s")
        plt.xlabel("Nombre de fichiers correspondants")
        plt.ylabel("Durée du test (secondes)")
        plt.title("Distribution des durées de test")
        plt.grid(True)
        plt.show()

    def update_sheet(self):
        if not self.sheet_available:
            self.log("tksheet manquant.")
            return
        sheet = self.sheet
        if sheet is None:
            return
        data_display = [[("" if v is None else v) for v in r] for r in self.pivot_rows]
        sheet.set_sheet_data(data_display)
        sheet.headers(self.pivot_headers)
        sheet.set_all_cell_sizes_to_text()
        try:
            etape_col_idx = self.pivot_headers.index("Étape")
            for r in range(len(self.pivot_rows)):
                sheet.highlight_cells(row=r, column=etape_col_idx, bg=COLOR_STEP_COLUMN_BG)
        except ValueError:
            pass
        self.recolor_table(initial=True)
        self.apply_column_visibility()
        self.filter_rows()

    def recolor_table(self, initial: bool = False):
        if not self.sheet_available or not self.pivot_rows:
            return
        sheet = self.sheet
        if sheet is None:
            return
        try:
            thr_near = float(self.var_thr_near.get())
            thr_yellow = float(self.var_thr_yellow.get())
        except ValueError:
            messagebox.showwarning("Seuils", "Valeurs seuils invalides.")
            return
        headers = self.pivot_headers
        if "Reference" not in headers:
            return
        idx_ref = headers.index("Reference")
        last_row_idx = len(self.pivot_rows) - 1
        use_seuils = thr_near > 0 or thr_yellow > 0
        for r_idx, row in enumerate(self.pivot_rows):
            if r_idx == last_row_idx:
                for c in range(len(headers)):
                    sheet.highlight_cells(row=r_idx, column=c, bg=COLOR_TOTAL_BG)
                continue
            ref_val = row[idx_ref]
            ref_num = ref_val if isinstance(ref_val, int) else None
            for c_idx, header in enumerate(headers):
                val = row[c_idx]
                if header in ("Index", "Étape"):
                    continue
                if header == "Reference":
                    sheet.highlight_cells(row=r_idx, column=c_idx, bg=COLOR_REF_BG)
                    continue
                if header.startswith("Δ_"):
                    if isinstance(val, int) and isinstance(ref_num, int) and ref_num > 0:
                        if use_seuils:
                            delta = val
                            pct = abs(delta) / ref_num * 100 if ref_num else 0
                            if delta == 0:
                                color = COLOR_EQUAL_DARK
                            elif pct <= thr_near:
                                color = COLOR_NEAR_LIGHT
                            elif pct <= thr_yellow:
                                color = COLOR_WARN
                            else:
                                color = COLOR_BAD
                        else:
                            color = COLOR_BAD if val > 0 else COLOR_EQUAL_DARK
                    else:
                        color = COLOR_DELTA_BG_DEFAULT
                    sheet.highlight_cells(row=r_idx, column=c_idx, bg=color)
                    continue
                if isinstance(val, int) and isinstance(ref_num, int) and ref_num > 0:
                    if use_seuils:
                        delta = val - ref_num
                        pct = abs(delta) / ref_num * 100 if ref_num else 0
                        if delta == 0:
                            color = COLOR_EQUAL_DARK
                        elif pct <= thr_near:
                            color = COLOR_NEAR_LIGHT
                        elif pct <= thr_yellow:
                            color = COLOR_WARN
                        else:
                            color = COLOR_BAD
                    else:
                        color = COLOR_BAD if val > ref_num else COLOR_EQUAL_DARK
                else:
                    color = COLOR_EMPTY
                sheet.highlight_cells(row=r_idx, column=c_idx, bg=color)
        if not initial:
            self.log("Recoloration appliquée.")

    def filter_rows(self):
        if not self.sheet_available or not self.pivot_rows:
            return
        sheet = self.sheet
        if sheet is None:
            return
        pattern = self.var_search.get().strip().lower()
        try:
            col_etape = self.pivot_headers.index("Étape")
        except ValueError:
            return
        total_idx = len(self.pivot_rows) - 1
        api_has_hide = hasattr(sheet, "hide_rows")
        api_has_show = hasattr(sheet, "show_rows")
        if api_has_hide and api_has_show:
            try:
                sheet.show_rows(iter(range(len(self.pivot_rows))))
            except Exception:
                pass
            if not pattern:
                return
            rows_to_hide = []
            for r_idx, row in enumerate(self.pivot_rows):
                if r_idx == total_idx:
                    continue
                if pattern not in str(row[col_etape]).lower():
                    rows_to_hide.append(r_idx)
            if rows_to_hide:
                try:
                    sheet.hide_rows(iter(rows_to_hide))
                except Exception:
                    pass
            return
        if not pattern:
            data_display = [[("" if v is None else v) for v in r] for r in self.pivot_rows]
            sheet.set_sheet_data(data_display)
            sheet.headers(self.pivot_headers)
            sheet.set_all_cell_sizes_to_text()
            self.recolor_table(initial=True)
            self.apply_column_visibility()
            return
        visible_data = []
        for i, r in enumerate(self.pivot_rows):
            if i == total_idx or pattern in str(r[col_etape]).lower():
                visible_data.append([("" if v is None else v) for v in r])
        sheet.set_sheet_data(visible_data)
        sheet.headers(self.pivot_headers)
        sheet.set_all_cell_sizes_to_text()
        self.recolor_table(initial=True)
        self.apply_column_visibility()

    def _auto_apply_search(self):
        self.filter_rows()
        self.after(600, self._auto_apply_search)

    def manage_columns(self):
        if not self.pivot_headers:
            messagebox.showinfo("Colonnes", "Construis d'abord le tableau.")
            return
        top = tk.Toplevel(self)
        top.title("Colonnes")
        top.geometry("420x480")
        frm = ttk.Frame(top, padding=8)
        frm.pack(fill="both", expand=True)
        canvas = tk.Canvas(frm)
        scroll_y = ttk.Scrollbar(frm, orient="vertical", command=canvas.yview)
        inner = ttk.Frame(canvas)
        inner.bind("<Configure>", lambda e: canvas.configure(scrollregion=canvas.bbox("all")))
        canvas.create_window((0, 0), window=inner, anchor="nw")
        canvas.configure(yscrollcommand=scroll_y.set)
        canvas.pack(side="left", fill="both", expand=True)
        scroll_y.pack(side="right", fill="y")
        vars_map: dict[int, tk.BooleanVar] = {}
        for idx, header in enumerate(self.pivot_headers):
            var = tk.BooleanVar(value=(idx not in self.hidden_cols))
            vars_map[idx] = var
            cb = ttk.Checkbutton(inner, text=f"{idx:02d}  {header}", variable=var)
            cb.pack(anchor="w", pady=2)
            if header in MANDATORY_COLUMNS:
                cb.state(["disabled"])
        def apply():
            new_hidden = {i for i, v in vars_map.items()
                          if not v.get() and self.pivot_headers[i] not in MANDATORY_COLUMNS}
            self.hidden_cols = new_hidden
            self.apply_column_visibility()
            top.destroy()
            self.log(f"Colonnes masquées: {sorted(self.hidden_cols)}")
        ttk.Button(top, text="Appliquer", command=apply).pack(side="bottom", pady=6)

    def apply_column_visibility(self):
        if not self.sheet_available or not self.pivot_headers:
            return
        sheet = self.sheet
        if sheet is None:
            return
        total_cols = len(self.pivot_headers)
        if hasattr(sheet, "show_columns"):
            try:
                sheet.show_columns(iter(range(total_cols)))
            except Exception:
                pass
        if hasattr(sheet, "hide_columns") and self.hidden_cols:
            try:
                sheet.hide_columns(iter(sorted(self.hidden_cols)))
            except Exception:
                pass

    def _save_prefs_state(self):
        data = {
            "ref_path": self.entry_ref.get().strip(),
            "logs_folder": self.entry_logs.get().strip(),
            "output_folder": self.entry_out.get().strip(),
            "show_deltas": self.var_show_deltas.get(),
            "thr_near": self.var_thr_near.get(),
            "thr_yellow": self.var_thr_yellow.get()
        }
        save_prefs(data)

    def destroy(self):
        self._save_prefs_state()
        super().destroy()

    def export_excel(self):
        if not self.pivot_headers or not self.pivot_rows:
            self.build_table()
            if not self.pivot_headers:
                return
        if not self.openpyxl_available:
            messagebox.showwarning("Excel", "openpyxl non installé (pip install openpyxl).")
            return
        try:
            if self.output_folder:
                out_dir = self.output_folder
            else:
                out_dir = self.reference_path.parent if self.reference_path else Path.cwd()
            out_dir.mkdir(parents=True, exist_ok=True)
            xlsx_path = out_dir / ((self.reference_path.stem if self.reference_path else "pivot") + ".comparaison_pivot_unique.xlsx")
            wb = openpyxl.Workbook()
            if SIGNATURE_ENABLED:
                wb.properties.creator = SIGNATURE_TEXT
                wb.properties.lastModifiedBy = SIGNATURE_TEXT
                wb.properties.comments = SIGNATURE_TEXT
            ws: _Any = wb.active
            ws.title = "Comparatif"
            def mk(col): return PatternFill("solid", fgColor=col.lstrip("#"))
            fill_equal = mk(COLOR_EQUAL_DARK)
            fill_near = mk(COLOR_NEAR_LIGHT)
            fill_warn = mk(COLOR_WARN)
            fill_bad = mk(COLOR_BAD)
            fill_empty = mk(COLOR_EMPTY)
            fill_total = mk(COLOR_TOTAL_BG)
            fill_step = mk(COLOR_STEP_COLUMN_BG)
            fill_ref = mk(COLOR_REF_BG)
            fill_delta_def = mk(COLOR_DELTA_BG_DEFAULT)
            bold = Font(bold=True)
            ws.append(self.pivot_headers)
            for c in range(1, len(self.pivot_headers) + 1):
                cell = ws.cell(row=1, column=c)
                cell.font = bold
                cell.alignment = Alignment(horizontal="center", vertical="center")
            if self.file_label_map_excel:
                for c in range(1, len(self.pivot_headers) + 1):
                    orig = self.pivot_headers[c - 1]
                    new_name = None
                    if orig in self.file_label_map_excel:
                        new_name = self.file_label_map_excel[orig]
                    elif orig.startswith("Δ_"):
                        base = orig[2:]
                        if base in self.file_label_map_excel:
                            new_name = f"Δ_{self.file_label_map_excel[base]}"
                    if new_name:
                        ws.cell(row=1, column=c).value = new_name
            try:
                idx_ref = self.pivot_headers.index("Reference")
            except ValueError:
                idx_ref = None
            thr_near = float(self.var_thr_near.get())
            thr_yellow = float(self.var_thr_yellow.get())
            last_row_index = len(self.pivot_rows) - 1
            use_seuils = thr_near > 0 or thr_yellow > 0
            def classify(ref_val: Optional[int], val: Any, is_delta: bool = False):
                if is_delta:
                    if isinstance(val, int) and isinstance(ref_val, int) and ref_val > 0:
                        if use_seuils:
                            delta = val
                            pct = abs(delta) / ref_val * 100 if ref_val else 0
                            if delta == 0:
                                return fill_equal
                            elif pct <= thr_near:
                                return fill_near
                            elif pct <= thr_yellow:
                                return fill_warn
                            else:
                                return fill_bad
                        else:
                            return fill_bad if val > 0 else fill_equal
                    return fill_delta_def
                else:
                    if isinstance(val, int) and isinstance(ref_val, int) and ref_val > 0:
                        if use_seuils:
                            delta = val - ref_val
                            pct = abs(delta) / ref_val * 100 if ref_val else 0
                            if delta == 0:
                                return fill_equal
                            elif pct <= thr_near:
                                return fill_near
                            elif pct <= thr_yellow:
                                return fill_warn
                            else:
                                return fill_bad
                        else:
                            return fill_bad if val > ref_val else fill_equal
                    return fill_empty
            for r_idx, row in enumerate(self.pivot_rows):
                ws.append(row)
                excel_r = r_idx + 2
                if r_idx == last_row_index:
                    for c in range(1, len(self.pivot_headers) + 1):
                        cell = ws.cell(row=excel_r, column=c)
                        cell.fill = fill_total
                        cell.font = bold
                    continue
                ref_val = row[idx_ref] if idx_ref is not None else None
                for c_idx, header in enumerate(self.pivot_headers):
                    cell = ws.cell(row=excel_r, column=c_idx + 1)
                    val = row[c_idx]
                    if header == "Étape":
                        cell.fill = fill_step
                        continue
                    if header == "Reference":
                        cell.fill = fill_ref
                        continue
                    if header == "Index":
                        continue
                    if header.startswith("Δ_"):
                        cell.fill = classify(ref_val, val, is_delta=True)
                        continue
                    cell.fill = classify(ref_val, val, is_delta=False)
            for col in range(1, len(self.pivot_headers) + 1):
                ws.column_dimensions[get_column_letter(col)].width = 14
            try:
                col_etape = self.pivot_headers.index("Étape") + 1
                ws.freeze_panes = f"{_gcl(col_etape + 1)}2"
            except ValueError:
                pass
            try:
                logs_dir = self.entry_logs.get().strip()
                if logs_dir:
                    folder = Path(logs_dir)
                elif self.reference_path:
                    folder = self.reference_path.parent
                else:
                    folder = None
                durees: List[int] = []
                if folder and folder.is_dir():
                    files = sorted(folder.glob("*.log"))
                    for p in files:
                        val = extract_duree_test_seconds(p)
                        if isinstance(val, int):
                            durees.append(val)
                if durees:
                    counts = Counter(durees)
                    xs = sorted(counts.keys())
                    ys = [counts[x] for x in xs]
                    ws_curve = wb.create_sheet(title="Courbe durée")
                    ws_curve.append(["duree_secondes", "nb_fichiers"])
                    for x, y in zip(xs, ys):
                        ws_curve.append([x, y])
                    chart = ScatterChart()
                    chart.scatterStyle = "lineMarker"
                    chart.title = "Distribution de la durée du test"
                    chart.style = 2
                    chart.x_axis.title = "Durée du test (secondes)"
                    chart.y_axis.title = "Nombre de fichiers"
                    chart.x_axis.number_format = "0"
                    chart.y_axis.number_format = "0"
                    chart.legend = None
                    max_row = 1 + len(xs)
                    xvalues = Reference(ws_curve, min_col=1, min_row=2, max_row=max_row)
                    yvalues = Reference(ws_curve, min_col=2, min_row=2, max_row=max_row)
                    series = Series(yvalues, xvalues, title="") # pyright: ignore[reportCallIssue]
                    chart.series.append(series)
                    chart.y_axis.majorGridlines = ChartLines()
                    chart.x_axis.majorGridlines = ChartLines()
                    try:
                        x_min, x_max = xs[0], xs[-1]
                        x_range = max(0, x_max - x_min)
                        x_unit = max(1, (x_range + 9) // 10) if x_range > 0 else 1
                        if not isinstance(chart.x_axis, TextAxis):
                            chart.x_axis.majorUnit = x_unit
                        chart.x_axis.scaling.min = max(0, x_min - x_unit)
                        chart.x_axis.scaling.max = x_max + x_unit
                        y_max = max(ys)
                        y_unit = max(1, (y_max + 9) // 10)
                        if not isinstance(chart.y_axis, TextAxis):
                            chart.y_axis.majorUnit = y_unit
                        chart.y_axis.scaling.min = 0
                        chart.y_axis.scaling.max = y_max + y_unit
                    except Exception:
                        pass
                    ws_curve.add_chart(chart, "E2")
            except Exception:
                self.log("Avertissement: impossible d'ajouter la feuille 'Courbe durée'.")
                self.log(traceback.format_exc())
            tmp_path = xlsx_path.with_suffix(".tmp.xlsx")
            try:
                wb.save(tmp_path)
                if xlsx_path.exists():
                    try:
                        xlsx_path.unlink()
                    except PermissionError:
                        alt = xlsx_path.parent / (xlsx_path.stem + f"_old_{int(time.time())}.xlsx")
                        try:
                            xlsx_path.rename(alt)
                        except Exception:
                            pass
                Path(tmp_path).rename(xlsx_path)
            except PermissionError:
                fallback = Path.home() / "logs_export"
                fallback.mkdir(parents=True, exist_ok=True)
                xlsx_path = fallback / xlsx_path.name
                wb.save(xlsx_path)
            self.log(f"Extraction Excel terminée : {xlsx_path}")
            messagebox.showinfo("Extraction Excel", f"Fichier créé :\n{xlsx_path}")
        except Exception:
            self.log("Erreur export Excel:\n" + traceback.format_exc())
            messagebox.showerror("Erreur Excel", "Impossible d'exporter vers Excel.")

def main():
    app = LogCompareApp()
    app.entry_ref.delete(0, "end")
    app.entry_logs.delete(0, "end")
    app.entry_out.delete(0, "end")
    app.mainloop()